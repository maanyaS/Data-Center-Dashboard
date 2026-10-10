"""Print the first rows of every sheet in the NYISO queue file + rows that look like loads.

Run from backend/:  python -m scripts.inspect_queue data/NYISO-Interconnection-Queue.xlsx
"""
import sys

from openpyxl import load_workbook

wb = load_workbook(sys.argv[1], read_only=True, data_only=True)
for ws in wb.worksheets:
    print(f"\n== SHEET: {ws.title}")
    loads = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        cells = ["" if c is None else str(c).strip() for c in row]
        if i < 6:
            print(f"row {i}: {cells}")
        elif "L" in cells or any("data center" in c.lower() for c in cells):
            loads.append(cells)
    print(f"-- {len(loads)} rows with an 'L' cell or 'data center'; first 3:")
    for r in loads[:3]:
        print("  ", r)
